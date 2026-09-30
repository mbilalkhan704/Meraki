"""
Meraki - utilities module.
Standalone helper functions with no dependency on the app's UI/state:
Google Fonts networking, tabular file reading/writing, styled-text
rendering onto certificate images, and color-picker math. Every function
here is pure/self-contained - none of them touch a CertificateApp instance.
"""

import os
import csv
import re
import json
import colorsys
import urllib.request
import urllib.parse
import urllib.error
from PIL import Image, ImageDraw # type: ignore

from mck_constants import (
    GOOGLE_FONTS_CSS_URL,
    _LEGACY_USER_AGENT,
    GOOGLE_FONTS_API_URL
)


# ---------------------------------------------------------------------------
# Google Fonts networking
# ---------------------------------------------------------------------------
def fetch_google_font_bytes(family_name: str) -> bytes:
    """Fallback path (no API key): scrape a direct .ttf link out of
    Google's font CSS endpoint. Raises on any failure."""
    query = urllib.parse.quote(family_name)
    css_url = GOOGLE_FONTS_CSS_URL.format(family=query)
    req = urllib.request.Request(css_url, headers={"User-Agent": _LEGACY_USER_AGENT})
    with urllib.request.urlopen(req, timeout=15) as resp:
        css_text = resp.read().decode("utf-8", errors="ignore")

    match = re.search(r"url\((https://fonts\.gstatic\.com/[^)]+?\.ttf)\)", css_text)
    if not match:
        raise RuntimeError(f"Google Fonts didn't return a .ttf file for '{family_name}'.")
    with urllib.request.urlopen(match.group(1), timeout=15) as resp:
        return resp.read()


def fetch_google_fonts_catalog(api_key: str) -> dict:
    """Return {family_name: direct_ttf_url} for every font Google Fonts
    offers, via the official documented API. Raises on a bad key / no
    internet. Used for the passive, already-trusted refresh at launch;
    for anything the user just typed in and is submitting for the first
    time, use validate_google_fonts_api_key instead - it distinguishes a
    genuinely bad key from a transient network problem, which this
    function does not."""
    url = f"{GOOGLE_FONTS_API_URL}?key={urllib.parse.quote(api_key)}"
    req = urllib.request.Request(url, headers={"User-Agent": "certificate-generator"})
    with urllib.request.urlopen(req, timeout=20) as resp:
        data = json.load(resp)

    catalog = {}
    for item in data.get("items", []):
        family = item.get("family")
        files = item.get("files", {})
        ttf_url = files.get("regular") or next(iter(files.values()), None)
        if family and ttf_url:
            catalog[family] = ttf_url.replace("http://", "https://")
    if not catalog:
        raise RuntimeError("Google returned no fonts - double check the API key.")
    return catalog


def validate_google_fonts_api_key(api_key: str, timeout: int = 12):
    """Actually check a key against Google's API instead of just trusting
    whatever the user typed - the only reliable way to know if a key
    really works. Returns a (status, payload) tuple:
      ("valid", catalog_dict)   - key works; payload is the font catalog
      ("invalid", message)      - Google rejected the key itself
      ("error", message)        - couldn't tell either way (no internet,
                                   timeout, Google-side outage, etc.) -
                                   NOT the same as invalid, and should
                                   never cause a key to be discarded.
    """
    url = f"{GOOGLE_FONTS_API_URL}?key={urllib.parse.quote(api_key)}"
    req = urllib.request.Request(url, headers={"User-Agent": "certificate-generator"})

    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.load(resp)
    except urllib.error.HTTPError as exc:
        try:
            body = json.loads(exc.read().decode("utf-8", errors="ignore"))
            message = body.get("error", {}).get("message", str(exc))
        except Exception:
            message = f"HTTP {exc.code}"
        if exc.code in (400, 401, 403):
            return "invalid", message
        return "error", message
    except urllib.error.URLError as exc:
        return "error", str(getattr(exc, "reason", exc))
    except Exception as exc:
        return "error", str(exc)

    catalog = {}
    for item in data.get("items", []):
        family = item.get("family")
        files = item.get("files", {})
        ttf_url = files.get("regular") or next(iter(files.values()), None)
        if family and ttf_url:
            catalog[family] = ttf_url.replace("http://", "https://")

    if not catalog:
        return "invalid", "Google didn't return any fonts for this key."
    return "valid", catalog


def fetch_bytes(url: str) -> bytes:
    with urllib.request.urlopen(url, timeout=15) as resp:
        return resp.read()


# ---------------------------------------------------------------------------
# Filenames
# ---------------------------------------------------------------------------
def sanitize_filename(name: str) -> str:
    name = name.strip()
    name = re.sub(r'[\\/*?:"<>|]', "", name)
    name = re.sub(r"\s+", " ", name)
    return name or "unnamed"


def name_to_filename(raw_name: str) -> str:
    """Turn a raw CSV/Excel name into a filename stem: each word
    capitalized, joined with underscores. e.g. 'bilal' -> 'Bilal',
    'Muhammad bilal khan' -> 'Muhammad_Bilal_Khan'."""
    cleaned = sanitize_filename(raw_name)
    words = cleaned.split(" ")
    words = [w.capitalize() for w in words if w]
    return "_".join(words) or "Unnamed"


# ---------------------------------------------------------------------------
# Tabular (CSV/Excel) file I/O
# ---------------------------------------------------------------------------
def read_tabular_file(path: str):
    """Read a .csv or .xlsx/.xlsm file and return (headers, rows) where
    rows is a list of dicts keyed by header. Raises on any problem."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            rows = [dict(r) for r in reader]
        return headers, rows

    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook # type: ignore
        except ImportError:
            raise RuntimeError(
                "Reading Excel files needs the 'openpyxl' package.\n"
                "Install it with:  pip install openpyxl"
            )
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], []
        headers = [str(h).strip() if h is not None else f"Column {i + 1}"
                   for i, h in enumerate(header_row)]
        rows = []
        for raw_row in rows_iter:
            if raw_row is None or all(v is None for v in raw_row):
                continue
            row_dict = {}
            for i, header in enumerate(headers):
                val = raw_row[i] if i < len(raw_row) else None
                row_dict[header] = "" if val is None else str(val)
            rows.append(row_dict)
        return headers, rows

    raise ValueError(f"Unsupported file type: {ext or '(none)'}. Use .csv or .xlsx.")


def write_tabular_file(path: str, headers: list, rows: list, new_columns: set):
    """Write (headers, rows) back to the original .csv/.xlsx/.xlsm file.

    `new_columns` lists which of `headers` are newly-added (e.g. a fresh
    Certificate ID column) - for Excel files, only cells in those columns
    get touched, so any existing formatting/content elsewhere in the
    sheet is left alone. CSV has no such concern since the whole file is
    just plain rows.

    Raises PermissionError (or another OSError) if the file is locked -
    e.g. currently open in Excel - so the caller can prompt the user to
    close it and retry, rather than silently failing or corrupting data.
    """
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for row in rows:
                writer.writerow({h: row.get(h, "") for h in headers})
        return

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook # type: ignore
        wb = load_workbook(path)  # NOT read_only - we need to write
        ws = wb.active

        existing_header_cells = list(ws[1])
        header_to_col = {}
        for cell in existing_header_cells:
            if cell.value is not None:
                header_to_col[str(cell.value).strip()] = cell.column

        next_col = len(existing_header_cells) + 1
        for h in headers:
            if h not in header_to_col:
                ws.cell(row=1, column=next_col, value=h)
                header_to_col[h] = next_col
                next_col += 1

        for r_idx, row in enumerate(rows, start=2):
            for h in new_columns:
                col = header_to_col[h]
                ws.cell(row=r_idx, column=col, value=row.get(h, ""))

        wb.save(path)
        return

    raise ValueError(f"Unsupported file type: {ext or '(none)'}. Use .csv or .xlsx.")


# ---------------------------------------------------------------------------
# Bold / italic / underline text rendering. Works with ANY font (no need
# for the font to have separate bold/italic files) using well-known faux
# techniques: a stroke outline for bold, an affine shear for italic, and a
# drawn line for underline.
# ---------------------------------------------------------------------------
def render_styled_text_layer(text, font, color, font_size, bold, italic, underline):
    """Return a small RGBA image containing just the styled text, sized to
    fit it tightly (plus a little padding for the bold stroke)."""
    stroke_w = max(1, round(font_size * 0.045)) if bold else 0

    scratch = Image.new("RGBA", (4, 4), (0, 0, 0, 0))
    scratch_draw = ImageDraw.Draw(scratch)
    bbox = scratch_draw.textbbox((0, 0), text, font=font, stroke_width=stroke_w)

    pad = stroke_w + 4
    w = max(1, (bbox[2] - bbox[0]) + pad * 2)
    h = max(1, (bbox[3] - bbox[1]) + pad * 2)

    layer = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    ld = ImageDraw.Draw(layer)
    draw_x = pad - bbox[0]
    draw_y = pad - bbox[1]

    if bold:
        ld.text((draw_x, draw_y), text, font=font, fill=color,
                 stroke_width=stroke_w, stroke_fill=color)
    else:
        ld.text((draw_x, draw_y), text, font=font, fill=color)

    if underline:
        try:
            ascent, _descent = font.getmetrics()
        except Exception:
            ascent = int(font_size * 0.8)
        line_w = max(1, round(font_size * 0.045))
        underline_y = draw_y + ascent + max(2, round(font_size * 0.06))
        text_w = bbox[2] - bbox[0]
        ld.line([(draw_x, underline_y), (draw_x + text_w, underline_y)],
                fill=color, width=line_w)

    if italic:
        shear = 0.22
        xshift = int(round(shear * h))
        new_w = w + xshift
        layer = layer.transform(
            (new_w, h), Image.AFFINE, (1, shear, 0, 0, 1, 0),
            resample=Image.BICUBIC,
        )

    return layer


def composite_styled_text(base_img, text, font, color, font_size, x, y, align,
                           bold, italic, underline):
    """Render styled text and paste it onto base_img so that (x, y) is the
    vertical center / chosen horizontal alignment anchor, matching the
    behavior of plain ImageDraw.text-based alignment used elsewhere."""
    layer = render_styled_text_layer(text, font, color, font_size, bold, italic, underline)
    lw, lh = layer.size
    if align == "center":
        dest_x = x - lw / 2
    elif align == "right":
        dest_x = x - lw
    else:
        dest_x = x
    dest_y = y - lh / 2
    base_img.paste(layer, (int(round(dest_x)), int(round(dest_y))), layer)


# ---------------------------------------------------------------------------
# Color-picker math (HSV spectrum/hue-strip image generation, hex helpers)
# ---------------------------------------------------------------------------
def _hls_hex(h, l, s):
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "#{:02x}{:02x}{:02x}".format(int(r * 255), int(g * 255), int(b * 255))


def _build_sv_square_image(hue, size=180):
    """A saturation(x) / value(y) square for a fixed hue, built from raw
    HSV byte planes - far faster than a per-pixel PIL draw loop."""
    h_byte = int(hue * 255)
    h_band = bytes([h_byte]) * (size * size)
    s_row = bytes([int(x / (size - 1) * 255) for x in range(size)])
    s_band = s_row * size
    v_band = bytes([int((1 - y / (size - 1)) * 255) for y in range(size) for _ in range(size)])
    img = Image.merge("HSV", (
        Image.frombytes("L", (size, size), h_band),
        Image.frombytes("L", (size, size), s_band),
        Image.frombytes("L", (size, size), v_band),
    ))
    return img.convert("RGB")


def _build_hue_strip_image(width=22, height=180):
    """A vertical strip covering every hue at full saturation/value."""
    h_band = bytes([int(y / (height - 1) * 255) for y in range(height) for _ in range(width)])
    s_band = bytes([255]) * (width * height)
    v_band = bytes([255]) * (width * height)
    img = Image.merge("HSV", (
        Image.frombytes("L", (width, height), h_band),
        Image.frombytes("L", (width, height), s_band),
        Image.frombytes("L", (width, height), v_band),
    ))
    return img.convert("RGB")
